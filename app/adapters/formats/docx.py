"""DOCX body and table paragraph extraction and translation rendering."""

import asyncio
import hashlib
import uuid
from pathlib import Path
from typing import Any

import structlog
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.errors import DocumentError, ErrorCode
from app.core.models import Block, DocumentIR, RenderResult

_logger = structlog.get_logger(__name__)


class DocxExtractor:
    """Extract body and top-level table paragraphs in document reading order."""

    async def extract(self, file_path: Path, document_id: str) -> DocumentIR:
        try:
            return await asyncio.to_thread(_extract_sync, file_path, document_id)
        except Exception:
            raise DocumentError(ErrorCode.CORRUPT_FILE) from None


class DocxRenderer:
    """Replace translated paragraph content on the original DOCX canvas."""

    async def render(
        self,
        original_path: Path,
        blocks: list[Block],
        translations: dict[str, str],
        output_path: Path,
    ) -> RenderResult:
        try:
            rendered_path = await asyncio.to_thread(
                _render_sync,
                original_path,
                blocks,
                translations,
                output_path,
            )
            return RenderResult(output_path=rendered_path)
        except Exception as error:
            _logger.error(
                "docx_render_failed", stage="docx_render", error_type=type(error).__name__
            )
            raise DocumentError(ErrorCode.RENDER_FAILED) from None


def _extract_sync(file_path: Path, document_id: str) -> DocumentIR:
    document = Document(str(file_path))
    blocks: list[Block] = []

    table_index = 0
    for body_index, child in enumerate(document.element.body.iterchildren()):
        if child.tag == qn("w:p"):
            _append_paragraph_block(
                blocks,
                document_id,
                Paragraph(child, document),
                {"container": "body", "body_index": body_index},
            )
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            # Keep XML references alive for the whole table: vertical merges can
            # alias cells from earlier rows, and discarded proxies can reuse ids.
            seen_cells: dict[int, Any] = {}
            for row_index, row in enumerate(table.rows):
                for cell_index, cell in enumerate(row.cells):
                    cell_id = id(cell._tc)
                    if cell_id in seen_cells:
                        continue
                    seen_cells[cell_id] = cell._tc
                    # cell.paragraphs contains direct paragraphs only, so nested
                    # table text is intentionally outside this adapter's scope.
                    for paragraph_index, paragraph in enumerate(cell.paragraphs):
                        _append_paragraph_block(
                            blocks,
                            document_id,
                            paragraph,
                            {
                                "container": "table",
                                "table_index": table_index,
                                "row_index": row_index,
                                "cell_index": cell_index,
                                "paragraph_index": paragraph_index,
                            },
                        )
            table_index += 1

    return DocumentIR(
        id=document_id,
        filename=file_path.name,
        format="docx",
        size_bytes=file_path.stat().st_size,
        page_count=None,
        blocks=blocks,
    )


def _append_paragraph_block(
    blocks: list[Block],
    document_id: str,
    paragraph: Paragraph,
    metadata: dict[str, Any],
) -> None:
    source_text = paragraph.text.strip()
    if not source_text:
        return

    seq = len(blocks)
    style = paragraph.style
    metadata["style"] = style.name if style is not None else "Normal"
    blocks.append(
        Block(
            id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"{document_id}:{seq}")),
            seq=seq,
            source_text=source_text,
            source_hash=hashlib.sha256(source_text.encode("utf-8")).hexdigest(),
            format_metadata=metadata,
        )
    )


def _render_sync(
    original_path: Path,
    blocks: list[Block],
    translations: dict[str, str],
    output_path: Path,
) -> Path:
    document = Document(str(original_path))
    replacements: dict[int, tuple[Paragraph, str]] = {}

    for block in blocks:
        if block.id not in translations:
            continue

        paragraph = _resolve_paragraph(document, block.format_metadata)
        paragraph_id = id(paragraph._p)
        if paragraph_id in replacements:
            raise ValueError("multiple blocks refer to one paragraph")

        translated_text = translations[block.id]
        if not isinstance(translated_text, str):
            raise ValueError("translation must be text")
        # Retain paragraph proxies, including their XML elements, until save.
        replacements[paragraph_id] = (paragraph, translated_text)

    for paragraph, translated_text in replacements.values():
        paragraph.clear()
        paragraph.add_run(translated_text)

    document.save(str(output_path))
    return output_path


def _index(metadata: dict[str, Any], key: str) -> int:
    value = metadata.get(key)
    if not isinstance(value, int) or isinstance(value, bool) or value < 0:
        raise ValueError("invalid paragraph locator metadata")
    return value


def _resolve_paragraph(document: DocxDocument, metadata: dict[str, Any]) -> Paragraph:
    style = metadata.get("style")
    if not isinstance(style, str):
        raise ValueError("invalid paragraph style metadata")

    # Historic locators address document.paragraphs, not body XML children.
    if "container" not in metadata:
        return document.paragraphs[_index(metadata, "paragraph_index")]

    container = metadata["container"]
    if container == "body":
        child = document.element.body[_index(metadata, "body_index")]
        if child.tag != qn("w:p"):
            raise ValueError("body locator does not refer to a paragraph")
        return Paragraph(child, document)
    if container == "table":
        table = document.tables[_index(metadata, "table_index")]
        row = table.rows[_index(metadata, "row_index")]
        cell = row.cells[_index(metadata, "cell_index")]
        return cell.paragraphs[_index(metadata, "paragraph_index")]
    raise ValueError("invalid paragraph container metadata")
