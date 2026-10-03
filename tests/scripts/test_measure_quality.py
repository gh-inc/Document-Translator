from __future__ import annotations

import json
from argparse import Namespace
from pathlib import Path

import pytest
from docx import Document

from app.adapters.llm import fake_provider
from app.config import Settings
from scripts import measure_quality
from scripts.measure_quality import (
    MeasurementError,
    chrf_score,
    preservation_metrics,
    require_live_provider,
)


def test_chrf_is_whitespace_insensitive_and_uses_character_orders() -> None:
    assert chrf_score("ab c", "a bc") == 100.0
    # Per-order P/R are averaged, then combined with beta=2. This fixture is
    # 38.888... under the standard effective-order chrF variant.
    assert chrf_score("abc", "abd") == pytest.approx(350 / 9)


def test_chrf_empty_and_nonmatching_inputs() -> None:
    assert chrf_score("", "") == 0.0
    assert chrf_score("", "hello") == 0.0
    assert chrf_score("hello", "") == 0.0


def test_number_and_placeholder_preservation_counts_duplicates_and_categories() -> None:
    source = "Order {{id}} on 2024-12-31 costs $12.50 and $12.50; %s [NAME]"
    translated = "Auftrag {{id}} am 2024-12-31 kostet $12.50; %s [NAME]"

    result = preservation_metrics(source, translated)

    assert result["placeholder"] == {
        "source_count": 3,
        "preserved_count": 3,
        "preservation_percent": 100.0,
    }
    assert result["date"] == {
        "source_count": 1,
        "preserved_count": 1,
        "preservation_percent": 100.0,
    }
    assert result["currency"] == {
        "source_count": 2,
        "preserved_count": 1,
        "preservation_percent": 50.0,
    }
    assert result["number"] == {
        "source_count": 0,
        "preserved_count": 0,
        "preservation_percent": None,
    }
    assert result["preservation_percent"] == pytest.approx(100 * 5 / 6)


def test_live_measurement_rejects_fake_provider_and_missing_key() -> None:
    with pytest.raises(MeasurementError, match="fake-provider runs are rejected"):
        require_live_provider(Settings(llm_provider="fake", openai_api_key="sk-not-used"))
    with pytest.raises(MeasurementError, match="requires OPENAI_API_KEY"):
        require_live_provider(Settings(llm_provider="openai", openai_api_key=""))


def test_live_measurement_accepts_configured_openai_without_revealing_key() -> None:
    settings = Settings(llm_provider="openai", openai_api_key="sk-test-secret")

    require_live_provider(settings)

    assert "sk-test-secret" not in repr(settings)


@pytest.mark.parametrize("with_reference", [False, True])
async def test_offline_measurement_runs_upload_triage_worker_render_and_metrics(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, with_reference: bool
) -> None:
    source_text = "Invoice {{id}} dated 2024-12-31 costs $12.50. Call %s [NAME]."
    document_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph(source_text)
    document.save(document_path)
    reference_path = tmp_path / "reference.txt"
    reference_path.write_text(source_text, encoding="utf-8")
    monkeypatch.setattr(measure_quality, "require_live_provider", lambda _settings: None)
    monkeypatch.setattr(fake_provider, "_count_usage_sync", lambda _input, _output: (10, 4))
    settings = Settings(
        llm_provider="fake",
        fake_latency_ms=1,
        database_path=tmp_path / "unused.db",
        upload_storage_path=tmp_path / "unused-uploads",
        output_storage_path=tmp_path / "unused-out",
    )
    args = Namespace(
        document=document_path,
        target_language="de",
        reference=reference_path if with_reference else None,
        timeout=10.0,
    )

    report = await measure_quality.measure(args, settings)

    assert report["quality"]["mode"] == ("reference" if with_reference else "back_translation")
    assert report["quality"]["metric"] == "chrF"
    assert report["quality"]["score"] > 0
    assert report["number_placeholder_preservation"]["preservation_percent"] == 100.0
    assert report["usage"]["job_count"] == (1 if with_reference else 2)
    assert report["jobs"]
    assert all(job["attempt_count"] >= 1 for job in report["jobs"])
    assert all(job["status"] == "done" for job in report["jobs"])
    assert "DOCX tables" in report["measurement_exclusions"][-1]


async def test_model_matrix_runs_isolated_pipelines_and_reports_known_cost(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    document_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Invoice {{id}} dated 2024-12-31 costs $12.50.")
    document.save(document_path)
    reference_path = tmp_path / "reference.txt"
    reference_path.write_text("Invoice {{id}} dated 2024-12-31 costs $12.50.", encoding="utf-8")
    guarded: list[str] = []
    database_paths: list[Path] = []
    real_create_app = measure_quality.create_app

    def record_app(settings: Settings, *, frontend_dir: Path) -> object:
        database_paths.append(settings.database_path)
        return real_create_app(settings, frontend_dir=frontend_dir)

    monkeypatch.setattr(measure_quality, "create_app", record_app)
    monkeypatch.setattr(
        measure_quality,
        "require_live_provider",
        lambda settings: guarded.append(settings.openai_model),
    )
    monkeypatch.setattr(fake_provider, "_count_usage_sync", lambda _input, _output: (10, 4))
    settings = Settings(llm_provider="fake", fake_latency_ms=1)
    args = Namespace(
        document=document_path,
        target_language="de",
        reference=reference_path,
        timeout=10.0,
        models="gpt-4o-mini,gpt-4o",
    )

    matrix = await measure_quality.measure_models(args, settings)

    assert guarded == ["gpt-4o-mini", "gpt-4o"]
    assert len(database_paths) == len(set(database_paths)) == 2
    reports = matrix["reports"]
    rows = matrix["comparison"]
    assert isinstance(reports, list)
    assert isinstance(rows, list)
    assert [row["model"] for row in rows] == guarded
    assert reports[0]["jobs"][0]["job_id"] != reports[1]["jobs"][0]["job_id"]
    for report, row in zip(reports, rows, strict=True):
        usage = report["pipeline_usage"]
        assert "triage usage is included separately" in report["measurement_exclusions"][0]
        assert usage["known_estimated_cost_usd"] == pytest.approx(
            usage["bulk_estimated_cost_usd"] + usage["triage_estimated_cost_usd"]
        )
        assert row["known_tokens_in"] == usage["known_tokens_in"]
        assert row["known_tokens_out"] == usage["known_tokens_out"]
        assert row["recorded_bulk_requests"] == usage["recorded_bulk_attempts"]
        assert usage["triage_request_count"] is None
        assert row["cost_per_million_mixed_tokens_usd"] == pytest.approx(
            row["known_estimated_run_cost_usd"]
            * 1_000_000
            / (row["known_tokens_in"] + row["known_tokens_out"])
        )


@pytest.mark.parametrize(
    "models", ["", ",", "gpt-4o-mini,", ",gpt-4o", "gpt-4o-mini,gpt-4o-mini", "unknown"]
)
async def test_model_list_is_rejected_before_a_run_or_cost(
    models: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    called = False

    async def fail_if_measured(*_args: object, **_kwargs: object) -> dict[str, object]:
        nonlocal called
        called = True
        raise AssertionError("measurement should not run")

    monkeypatch.setattr(measure_quality, "measure", fail_if_measured)
    with pytest.raises(MeasurementError, match="--models"):
        await measure_quality.measure_models(
            Namespace(models=models), Settings(llm_provider="fake")
        )
    assert not called


def test_human_report_and_json_payload_can_be_emitted_on_separate_streams(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    report = {
        "measurement": {
            "date_utc": "2026-10-03",
            "sample_filename": "sample.docx",
            "provider": "openai",
            "model": "gpt-4o-mini",
            "source_block_count": 1,
        },
        "quality": {"mode": "reference", "score": 71.25},
        "number_placeholder_preservation": {"preservation_percent": 100.0},
        "jobs": [],
        "usage": {
            "job_count": 1,
            "cost_usd": 0.01,
            "retry_cost_share_percent": 0.0,
            "tokens_in": 10,
            "tokens_out": 12,
            "attempt_count": 1,
            "p95_chunk_attempt_latency_ms": 20,
            "p95_job_latency_ms": None,
        },
    }
    monkeypatch.setattr(measure_quality, "build_parser", lambda: _ArgsParser())
    monkeypatch.setattr(measure_quality, "Settings", lambda **_kwargs: Settings())

    def return_report(awaitable: object) -> dict[str, object]:
        awaitable.close()  # type: ignore[attr-defined]
        return report

    monkeypatch.setattr(measure_quality.asyncio, "run", return_report)

    assert measure_quality.main() == 0
    output = capsys.readouterr()

    assert json.loads(output.out) == report
    assert "Metric" in output.err
    assert "71.25 / 100" in output.err


def test_matrix_cli_emits_json_and_markdown_with_truthful_request_label(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    matrix: dict[str, object] = {
        "comparison": [
            {
                "model": "gpt-4o-mini",
                "cost_per_million_mixed_tokens_usd": 0.25,
                "known_estimated_run_cost_usd": 0.00025,
                "known_tokens_in": 800,
                "known_tokens_out": 200,
                "chrf": 71.25,
                "preservation_percent": None,
                "recorded_bulk_requests": 3,
            }
        ],
        "reports": [{"quality": {"mode": "reference"}}],
        "notes": ["Triage request counts are not persisted."],
    }
    monkeypatch.setattr(measure_quality, "build_parser", lambda: _ArgsParser(models="gpt-4o-mini"))
    monkeypatch.setattr(measure_quality, "Settings", lambda **_kwargs: Settings())

    def return_matrix(awaitable: object) -> dict[str, object]:
        awaitable.close()  # type: ignore[attr-defined]
        return matrix

    monkeypatch.setattr(measure_quality.asyncio, "run", return_matrix)

    assert measure_quality.main() == 0
    output = capsys.readouterr()
    assert json.loads(output.out) == matrix
    assert "| Model | Cost / 1M tokens in+out (USD)" in output.err
    assert "Quality mode: reference" in output.err
    assert "Recorded bulk requests" in output.err
    assert "| gpt-4o-mini | $0.2500 | $0.000250 | 71.25 | n/a | 3 | 800 / 200 |" in output.err


def test_failed_matrix_emits_no_scores_or_json(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(measure_quality, "build_parser", lambda: _ArgsParser(models="gpt-4o"))
    monkeypatch.setattr(measure_quality, "Settings", lambda **_kwargs: Settings())

    def fail(awaitable: object) -> dict[str, object]:
        awaitable.close()  # type: ignore[attr-defined]
        raise MeasurementError("second model failed")

    monkeypatch.setattr(measure_quality.asyncio, "run", fail)

    assert measure_quality.main() == 1
    output = capsys.readouterr()
    assert output.out == ""
    assert "second model failed" in output.err
    assert "chrF" not in output.err


def test_cli_exits_nonzero_with_explicit_fake_provider_message(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(measure_quality, "build_parser", lambda: _ArgsParser())
    monkeypatch.setattr(
        measure_quality,
        "Settings",
        lambda **_kwargs: Settings(llm_provider="fake", openai_api_key="sk-not-used"),
    )

    assert measure_quality.main() == 1
    output = capsys.readouterr()

    assert output.out == ""
    assert "fake-provider runs are rejected" in output.err


class _ArgsParser:
    def __init__(self, models: str | None = None) -> None:
        self.models = models

    def parse_args(self) -> Namespace:
        return Namespace(
            env_file=None,
            document=Path("unused.docx"),
            target_language="de",
            reference=None,
            timeout=1.0,
            models=self.models,
        )
