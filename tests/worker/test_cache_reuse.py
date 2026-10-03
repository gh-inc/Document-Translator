"""Real DOCX extraction and worker execution reuse unchanged paragraphs."""

from __future__ import annotations

import asyncio
from pathlib import Path

from docx import Document

from app.adapters.formats.docx import DocxExtractor, DocxRenderer
from app.adapters.formats.registry import FormatRegistry
from app.adapters.llm.fake_provider import FakeProvider
from app.adapters.llm.pricing import ModelCostCalculator
from app.adapters.persistence.api import ApiJobExecutionRepository, ApiPersistence
from app.adapters.persistence.database import SqliteConnectionFactory, transaction
from app.adapters.persistence.repositories import (
    SqliteDocumentRepository,
    SqliteJobExecutionRepository,
    SqliteTranslationCacheRepository,
)
from app.adapters.persistence.worker import WorkerPersistence
from app.adapters.storage.filesystem import FilesystemStorage
from app.config import Settings
from app.core.models import DocumentStatus, JobStatus, TranslationPlan
from app.core.services.job_service import JobService
from app.worker.claim_loop import ClaimLoop
from tests.worker.test_worker_integration import _wait_for_terminal


class RecordingFakeProvider(FakeProvider):
    def __init__(self, settings: Settings) -> None:
        super().__init__(settings=settings, fail_rate=0, latency_ms=0)
        self.calls: list[str] = []

    async def translate_chunk(self, request):
        self.calls.extend(block.source_text for block in request.blocks)
        return await super().translate_chunk(request)


async def test_edited_docx_reuses_unchanged_paragraphs_and_isolates_plan(tmp_path: Path) -> None:
    settings = Settings(
        database_path=tmp_path / "cache.db",
        upload_storage_path=tmp_path / "uploads",
        output_storage_path=tmp_path / "outputs",
        worker_id="cache-worker",
        heartbeat_interval_seconds=1,
        job_lease_seconds=5,
        chunk_lease_seconds=5,
        max_chunk_concurrency=1,
    )
    connection = await SqliteConnectionFactory(settings.database_path).create()
    shutdown = asyncio.Event()
    task: asyncio.Task[None] | None = None
    try:
        storage = FilesystemStorage(settings.upload_storage_path, settings.output_storage_path)
        documents = SqliteDocumentRepository(connection)
        jobs = SqliteJobExecutionRepository(connection, worker_id=settings.worker_id)
        cache = SqliteTranslationCacheRepository(connection)
        persistence = WorkerPersistence(connection)
        formats = FormatRegistry()
        formats.register("docx", DocxExtractor(), DocxRenderer())
        provider = RecordingFakeProvider(settings)
        worker = ClaimLoop(
            settings,
            jobs,
            cache,
            documents,
            provider,
            ModelCostCalculator(),
            formats,
            storage,
            persistence=persistence,
            shutdown_event=shutdown,
        )
        service = JobService(
            documents,
            ApiJobExecutionRepository(connection),
            cache,
            ModelCostCalculator(),
            persistence=ApiPersistence(connection),
            settings=settings,
        )
        task = asyncio.create_task(worker.run())

        async def translate(document_id: str, paragraphs: list[str], *, domain: str = "general"):
            source = tmp_path / f"{document_id}.docx"
            document = Document()
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)
            await asyncio.to_thread(document.save, str(source))
            content = await asyncio.to_thread(source.read_bytes)
            upload_path = await storage.save_upload(document_id, content, source.name)
            extracted = await DocxExtractor().extract(upload_path, document_id)
            assert [block.source_text for block in extracted.blocks] == paragraphs
            async with transaction(connection):
                await documents.create_document(
                    document_id,
                    extracted.filename,
                    "docx",
                    extracted.size_bytes,
                    str(upload_path),
                    extracted.page_count,
                )
                await documents.create_blocks(document_id, extracted.blocks)
                await documents.update_document_status(document_id, DocumentStatus.EXTRACTED)
                await documents.save_analysis(
                    document_id,
                    TranslationPlan(source_language="en", domain=domain, register="neutral"),
                )
            job = (await service.create_jobs(document_id, ["de"], f"request-{document_id}"))[0]
            completed = await _wait_for_terminal(persistence, jobs, job.id)
            assert completed.status is JobStatus.DONE
            output_path = await storage.get_output_path(job.id)
            rendered = await asyncio.to_thread(
                lambda: [paragraph.text for paragraph in Document(str(output_path)).paragraphs]
            )
            assert rendered == [f"[de] {paragraph}" for paragraph in paragraphs]
            return completed

        original = ["Alpha paragraph.", "Beta paragraph.", "Gamma paragraph."]
        first = await translate("first", original)
        assert provider.calls == original
        assert (first.cache_hit_blocks, first.cache_miss_blocks) == (0, 3)

        revised = [original[0], "Beta paragraph, revised.", original[2]]
        second = await translate("second", revised)
        assert provider.calls == original + [revised[1]]
        assert (second.cache_hit_blocks, second.cache_miss_blocks) == (2, 1)
        assert round(100 * second.cache_hit_blocks / 3) == 67

        legal = await translate("legal", original, domain="legal")
        assert provider.calls == original + [revised[1]] + original
        assert (legal.cache_hit_blocks, legal.cache_miss_blocks) == (0, 3)
    finally:
        shutdown.set()
        if task is not None:
            task.cancel()
            await asyncio.gather(task, return_exceptions=True)
        await connection.close()
