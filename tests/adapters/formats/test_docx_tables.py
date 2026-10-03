"""DOCX table translation, locator compatibility, and canvas regressions."""

from copy import deepcopy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

from app.adapters.formats.docx import DocxExtractor, DocxRenderer
from app.adapters.llm.fake_provider import FakeProvider
from app.core.errors import DocumentError, ErrorCode
from app.core.models import Block, ChunkRequest, TranslationPlan

_PLATON_SAMPLE = Path(__file__).resolve().parents[3] / "samples" / "platon-complex.docx"


def _block(block_id: str, metadata: dict[str, Any]) -> Block:
    return Block(
        id=block_id,
        seq=0,
        source_text="PRIVATE SOURCE",
        source_hash="hash",
        format_metadata=metadata,
    )


def _paragraphs(document: DocxDocument) -> list[Paragraph]:
    """Read the original canvas independently of adapter-produced locators."""
    paragraphs: list[Paragraph] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            paragraphs.append(item)
        else:
            seen: set[Any] = set()
            for row in item.rows:
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        paragraphs.extend(cell.paragraphs)
    return paragraphs


def _canvas_structure(document: DocxDocument) -> bytes:
    """Retain properties, empty paragraphs, tables, and unsupported content."""
    body = deepcopy(document.element.body)
    for child in body:
        if child.tag == qn("w:p"):
            candidates = [child]
        elif child.tag == qn("w:tbl"):
            candidates = child.xpath("./w:tr/w:tc/w:p")
        else:
            candidates = []
        for paragraph in candidates:
            for content in list(paragraph):
                if content.tag != qn("w:pPr"):
                    paragraph.remove(content)
    return etree.tostring(body, method="c14n")


@pytest.mark.asyncio
async def test_interleaved_body_and_cell_paragraphs_keep_reading_order(tmp_path: Path) -> None:
    source = tmp_path / "interleaved.docx"
    document = Document()
    document.add_paragraph("Before", style="Heading 1")
    document.add_paragraph("   ")
    document.element.body.insert(2, OxmlElement("w:customXml"))
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First cell"
    table.cell(0, 0).add_paragraph()
    table.cell(0, 0).add_paragraph("Second paragraph", style="Body Text")
    document.add_paragraph("Between")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Second table"
    document.add_paragraph("After")
    document.save(str(source))

    blocks = (await DocxExtractor().extract(source, "interleaved")).blocks

    assert [block.source_text for block in blocks] == [
        "Before",
        "First cell",
        "Second paragraph",
        "Between",
        "Second table",
        "After",
    ]
    assert [block.seq for block in blocks] == list(range(6))
    assert [block.format_metadata for block in blocks] == [
        {"container": "body", "body_index": 0, "style": "Heading 1"},
        {
            "container": "table",
            "table_index": 0,
            "row_index": 0,
            "cell_index": 0,
            "paragraph_index": 0,
            "style": "Normal",
        },
        {
            "container": "table",
            "table_index": 0,
            "row_index": 0,
            "cell_index": 0,
            "paragraph_index": 2,
            "style": "Body Text",
        },
        {"container": "body", "body_index": 4, "style": "Normal"},
        {
            "container": "table",
            "table_index": 1,
            "row_index": 0,
            "cell_index": 0,
            "paragraph_index": 0,
            "style": "Normal",
        },
        {"container": "body", "body_index": 6, "style": "Normal"},
    ]
    output = tmp_path / "out.docx"
    await DocxRenderer().render(
        source, blocks, {block.id: f"Translated {block.seq}" for block in blocks}, output
    )
    translated = Document(str(output))
    assert [p.text for p in _paragraphs(translated) if p.text.strip()] == [
        f"Translated {seq}" for seq in range(6)
    ]
    assert _canvas_structure(translated) == _canvas_structure(document)


@pytest.mark.asyncio
async def test_cell_translation_flattens_runs_and_preserves_paragraph_properties(
    tmp_path: Path,
) -> None:
    source = tmp_path / "styled-cell.docx"
    document = Document()
    paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph.style = "Body Text"
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run("First ").bold = True
    paragraph.add_run("second").italic = True
    document.save(str(source))
    blocks = (await DocxExtractor().extract(source, "styled-cell")).blocks
    assert [block.source_text for block in blocks] == ["First second"]
    output = tmp_path / "out.docx"

    await DocxRenderer().render(source, blocks, {blocks[0].id: "Translated cell"}, output)

    rendered = Document(str(output)).tables[0].cell(0, 0).paragraphs[0]
    assert rendered.text == "Translated cell"
    assert rendered.style.name == "Body Text"
    assert rendered.paragraph_format.keep_with_next is True
    assert rendered.paragraph_format.keep_together is True
    assert [run.text for run in rendered.runs] == ["Translated cell"]
    assert rendered.runs[0].bold is None
    assert rendered.runs[0].italic is None


@pytest.mark.asyncio
async def test_horizontal_and_vertical_merged_cells_are_extracted_once(tmp_path: Path) -> None:
    source = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Horizontal"
    table.cell(0, 2).merge(table.cell(2, 2)).text = "Vertical"
    table.cell(1, 0).text = "Middle left"
    table.cell(1, 1).text = "Middle right"
    table.cell(2, 0).text = "Bottom left"
    table.cell(2, 1).text = "Bottom right"
    document.save(str(source))

    blocks = (await DocxExtractor().extract(source, "merged")).blocks

    assert [block.source_text for block in blocks] == [
        "Horizontal",
        "Vertical",
        "Middle left",
        "Middle right",
        "Bottom left",
        "Bottom right",
    ]
    output = tmp_path / "out.docx"
    await DocxRenderer().render(
        source, blocks, {block.id: f"Translated {block.seq}" for block in blocks}, output
    )
    rendered = Document(str(output))
    assert [p.text for p in _paragraphs(rendered)] == [f"Translated {seq}" for seq in range(6)]
    assert _canvas_structure(rendered) == _canvas_structure(document)


@pytest.mark.asyncio
async def test_table_rows_with_omitted_grid_cells_keep_resolvable_locators(tmp_path: Path) -> None:
    source = tmp_path / "omitted-cells.docx"
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(1, 1).text = "Only populated cell"
    row = table.rows[1]
    row._tr.remove(row._tr.tc_lst[-1])
    row._tr.remove(row._tr.tc_lst[0])
    properties = row._tr.get_or_add_trPr()
    for name in ("w:gridBefore", "w:gridAfter"):
        omitted = OxmlElement(name)
        omitted.set(qn("w:val"), "1")
        properties.append(omitted)
    document.save(str(source))

    blocks = (await DocxExtractor().extract(source, "omitted")).blocks

    assert [block.source_text for block in blocks] == ["Only populated cell"]
    assert blocks[0].format_metadata == {
        "container": "table",
        "table_index": 0,
        "row_index": 1,
        "cell_index": 0,
        "paragraph_index": 0,
        "style": "Normal",
    }
    output = tmp_path / "out.docx"
    await DocxRenderer().render(source, blocks, {blocks[0].id: "Translated cell"}, output)
    rendered = Document(str(output))
    rendered_row = rendered.tables[0].rows[1]
    assert len(rendered_row.cells) == 1
    assert rendered_row.cells[0].text == "Translated cell"
    assert rendered_row.grid_cols_before == rendered_row.grid_cols_after == 1
    assert _canvas_structure(rendered) == _canvas_structure(document)


@pytest.mark.asyncio
async def test_nested_tables_headers_and_footers_remain_unchanged(tmp_path: Path) -> None:
    source = tmp_path / "nested.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "Outer cell"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested stays"
    cell.add_paragraph("After nested")
    document.sections[0].header.paragraphs[0].text = "Header stays"
    document.sections[0].footer.paragraphs[0].text = "Footer stays"
    nested_xml = nested._tbl.xml
    header_xml = document.sections[0].header._element.xml
    footer_xml = document.sections[0].footer._element.xml
    document.save(str(source))

    blocks = (await DocxExtractor().extract(source, "nested")).blocks
    assert [block.source_text for block in blocks] == ["Outer cell", "After nested"]
    output = tmp_path / "out.docx"
    await DocxRenderer().render(
        source, blocks, {block.id: f"Translated {block.seq}" for block in blocks}, output
    )
    rendered = Document(str(output))
    assert rendered.tables[0].cell(0, 0).tables[0]._tbl.xml == nested_xml
    assert rendered.sections[0].header._element.xml == header_xml
    assert rendered.sections[0].footer._element.xml == footer_xml
    assert [p.text for p in rendered.tables[0].cell(0, 0).paragraphs] == [
        "Translated 0",
        "",
        "Translated 1",
    ]


@pytest.mark.asyncio
async def test_legacy_paragraph_index_after_table_keeps_original_meaning(tmp_path: Path) -> None:
    source = tmp_path / "legacy.docx"
    document = Document()
    document.add_paragraph("Before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table stays"
    document.add_paragraph("After", style="Heading 2")
    document.save(str(source))
    block = _block("legacy", {"paragraph_index": 1, "style": "Heading 2"})
    output = tmp_path / "out.docx"

    await DocxRenderer().render(source, [block], {block.id: "Translated after"}, output)

    rendered = Document(str(output))
    assert [p.text for p in rendered.paragraphs] == ["Before", "Translated after"]
    assert rendered.tables[0].cell(0, 0).text == "Table stays"
    assert rendered.paragraphs[1].style.name == "Heading 2"


_TABLE_LOCATOR = {
    "container": "table",
    "table_index": 0,
    "row_index": 0,
    "cell_index": 0,
    "paragraph_index": 0,
    "style": "Normal",
}


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "metadata",
    [
        {"container": None, "paragraph_index": 0, "style": "Normal"},
        {"container": "unknown", "paragraph_index": 0, "style": "Normal"},
        {"container": "body", "body_index": 1, "style": "Normal"},  # table, not paragraph
        {"container": "body", "body_index": 99, "style": "Normal"},
        {"container": "body", "body_index": True, "style": "Normal"},
        {"container": "body", "body_index": -1, "style": "Normal"},
        {"container": "body", "body_index": "0", "style": "Normal"},
        {"container": "body", "style": "Normal"},
        {"container": "body", "body_index": 0, "style": None},
        *[
            {**_TABLE_LOCATOR, key: value}
            for key in ("table_index", "row_index", "cell_index", "paragraph_index")
            for value in (True, -1, "0", 99, None)
        ],
    ],
)
async def test_invalid_locators_fail_safely_before_output_is_written(
    tmp_path: Path, metadata: dict[str, Any]
) -> None:
    source = tmp_path / "invalid.docx"
    document = Document()
    document.add_paragraph("PRIVATE SOURCE")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "PRIVATE CELL"
    document.save(str(source))
    original_bytes = source.read_bytes()
    block = _block("invalid", metadata)
    output = tmp_path / "out.docx"

    with pytest.raises(DocumentError) as error:
        await DocxRenderer().render(source, [block], {block.id: "PRIVATE TRANSLATION"}, output)

    assert error.value.error_code is ErrorCode.RENDER_FAILED
    assert "PRIVATE" not in str(error.value)
    assert not output.exists()
    assert source.read_bytes() == original_bytes


@pytest.mark.asyncio
@pytest.mark.parametrize("merge", ["horizontal", "vertical", "legacy_body"])
async def test_duplicate_physical_targets_are_rejected(tmp_path: Path, merge: str) -> None:
    source = tmp_path / "duplicate.docx"
    document = Document()
    document.add_paragraph("Body")
    table = document.add_table(rows=2, cols=2)
    if merge == "legacy_body":
        first_metadata = {"paragraph_index": 0, "style": "Normal"}
        second_metadata = {"container": "body", "body_index": 0, "style": "Normal"}
    else:
        end = table.cell(0, 1) if merge == "horizontal" else table.cell(1, 0)
        table.cell(0, 0).merge(end).text = "Merged"
        first_metadata = dict(_TABLE_LOCATOR)
        second_metadata = {
            **_TABLE_LOCATOR,
            "cell_index" if merge == "horizontal" else "row_index": 1,
        }
    document.save(str(source))
    blocks = [_block("first", first_metadata), _block("second", second_metadata)]
    output = tmp_path / "out.docx"

    with pytest.raises(DocumentError) as error:
        await DocxRenderer().render(source, blocks, {"first": "One", "second": "Two"}, output)

    assert error.value.error_code is ErrorCode.RENDER_FAILED
    assert not output.exists()


@pytest.mark.asyncio
async def test_platon_sample_translates_all_44_body_and_table_paragraphs(tmp_path: Path) -> None:
    source = _PLATON_SAMPLE
    original = Document(str(source))
    source_paragraphs = [paragraph for paragraph in _paragraphs(original) if paragraph.text.strip()]
    source_texts = [paragraph.text.strip() for paragraph in source_paragraphs]
    original_styles = [paragraph.style.name for paragraph in source_paragraphs]
    assert len([p for p in original.paragraphs if p.text.strip()]) == 16
    assert len(source_paragraphs) == 44

    ir = await DocxExtractor().extract(source, "platon-sample")
    assert [block.source_text for block in ir.blocks] == source_texts
    assert [block.seq for block in ir.blocks] == list(range(44))
    assert sum(block.format_metadata["container"] == "table" for block in ir.blocks) == 28
    result = await FakeProvider(fail_rate=0, latency_ms=0).translate_chunk(
        ChunkRequest(
            chunk_id="platon",
            blocks=ir.blocks,
            target_language="DE",
            model="gpt-4o-mini",
            plan=TranslationPlan(source_language="RU", domain="general", register="neutral"),
        )
    )
    output = tmp_path / "platon-translated.docx"
    rendered = await DocxRenderer().render(source, ir.blocks, result.translations, output)
    assert rendered.degraded_block_ids == []
    assert rendered.fallback_blocks == rendered.fallback_pages == 0
    translated = Document(str(output))
    translated_paragraphs = [p for p in _paragraphs(translated) if p.text.strip()]
    assert [p.text for p in translated_paragraphs] == [f"[DE] {text}" for text in source_texts]
    assert [p.style.name for p in translated_paragraphs] == original_styles
    assert _canvas_structure(translated) == _canvas_structure(original)
