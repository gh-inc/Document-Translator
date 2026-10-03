"""DOCX adapter behavior and thread-boundary tests."""

import hashlib
import threading
import uuid
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from structlog.testing import capture_logs

from app.adapters.formats import docx as docx_adapter
from app.core.errors import DocumentError, ErrorCode
from app.core.models import Block


def _add_hyperlink(paragraph: Paragraph, text: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _make_document(path: Path) -> None:
    document = Document()
    document.add_paragraph()
    heading = document.add_paragraph("Original heading", style="Heading 1")
    heading.paragraph_format.keep_with_next = True

    linked = document.add_paragraph(style="Body Text")
    linked.add_run("Read ")
    _add_hyperlink(linked, "this link")
    linked.add_run(" now")

    unchanged = document.add_paragraph()
    unchanged.add_run("Keep ")
    emphasized = unchanged.add_run("this")
    emphasized.bold = True

    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table stays"
    document.sections[0].header.paragraphs[0].text = "Header stays"
    document.save(str(path))


@pytest.mark.asyncio
async def test_extractor_returns_nonempty_paragraphs_in_order_with_stable_identity(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    _make_document(source)
    extractor = docx_adapter.DocxExtractor()

    first = await extractor.extract(source, "document-1")
    second = await extractor.extract(source, "document-1")

    assert first.format == "docx"
    assert first.filename == "source.docx"
    assert first.size_bytes == source.stat().st_size
    assert first.page_count is None
    assert first.warnings == []
    assert [block.seq for block in first.blocks] == [0, 1, 2, 3]
    assert [block.source_text for block in first.blocks] == [
        "Original heading",
        "Read this link now",
        "Keep this",
        "Table stays",
    ]
    assert [block.format_metadata["body_index"] for block in first.blocks[:3]] == [
        1,
        2,
        3,
    ]
    assert [block.format_metadata["style"] for block in first.blocks] == [
        "Heading 1",
        "Body Text",
        "Normal",
        "Normal",
    ]
    assert [block.format_metadata["container"] for block in first.blocks] == [
        "body",
        "body",
        "body",
        "table",
    ]
    assert [block.id for block in first.blocks] == [block.id for block in second.blocks]
    assert first.blocks[0].id == str(uuid.uuid5(uuid.NAMESPACE_URL, "document-1:0"))
    assert first.blocks[0].source_hash == hashlib.sha256(b"Original heading").hexdigest()
    assert all(
        block.source_hash == hashlib.sha256(block.source_text.encode()).hexdigest()
        for block in first.blocks
    )


@pytest.mark.asyncio
async def test_renderer_replaces_only_translated_top_level_paragraphs(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "translated.docx"
    _make_document(source)
    original_document = Document(str(source))
    original_heading = original_document.paragraphs[1].text
    blocks = (await docx_adapter.DocxExtractor().extract(source, "document-1")).blocks

    rendered = await docx_adapter.DocxRenderer().render(
        source,
        blocks,
        {
            blocks[0].id: "Translated heading",
            blocks[1].id: "Translated linked text",
        },
        output,
    )

    assert rendered.output_path == output
    assert rendered.degraded_block_ids == []
    assert rendered.fallback_blocks == rendered.fallback_pages == 0
    translated = Document(str(output))
    assert translated.paragraphs[1].text == "Translated heading"
    heading_style = translated.paragraphs[1].style
    assert heading_style is not None
    assert heading_style.name == "Heading 1"
    assert translated.paragraphs[1].paragraph_format.keep_with_next is True
    assert [run.text for run in translated.paragraphs[1].runs] == ["Translated heading"]

    linked = translated.paragraphs[2]
    assert linked.text == "Translated linked text"
    linked_style = linked.style
    assert linked_style is not None
    assert linked_style.name == "Body Text"
    assert len(linked.runs) == 1
    assert linked.runs[0].bold is None
    assert linked.hyperlinks == []

    unchanged = translated.paragraphs[3]
    assert unchanged.text == "Keep this"
    assert len(unchanged.runs) == 2
    assert unchanged.runs[1].bold is True
    assert translated.tables[0].cell(0, 0).text == "Table stays"
    assert translated.sections[0].header.paragraphs[0].text == "Header stays"
    assert Document(str(source)).paragraphs[1].text == original_heading


@pytest.mark.asyncio
async def test_extractor_and_renderer_run_python_docx_in_worker_threads(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "translated.docx"
    _make_document(source)
    main_thread = threading.get_ident()
    document_threads: list[int] = []
    original_document = Document

    def tracking_document(*args: Any, **kwargs: Any) -> Any:
        document_threads.append(threading.get_ident())
        return original_document(*args, **kwargs)

    monkeypatch.setattr(docx_adapter, "Document", tracking_document)
    blocks = (await docx_adapter.DocxExtractor().extract(source, "document-1")).blocks
    await docx_adapter.DocxRenderer().render(source, blocks, {}, output)

    assert len(document_threads) == 2
    assert all(thread_id != main_thread for thread_id in document_threads)


@pytest.mark.asyncio
async def test_corrupt_input_errors_are_catalogued_and_suppress_library_details(
    tmp_path: Path,
) -> None:
    corrupt = tmp_path / "broken.docx"
    corrupt.write_bytes(b"not a word document")

    with pytest.raises(DocumentError) as extraction_error:
        await docx_adapter.DocxExtractor().extract(corrupt, "document-1")
    assert extraction_error.value.error_code is ErrorCode.CORRUPT_FILE
    assert "PackageNotFoundError" not in str(extraction_error.value)

    with pytest.raises(DocumentError) as render_error:
        await docx_adapter.DocxRenderer().render(corrupt, [], {}, tmp_path / "out.docx")
    assert render_error.value.error_code is ErrorCode.RENDER_FAILED
    assert "PackageNotFoundError" not in str(render_error.value)


@pytest.mark.asyncio
async def test_renderer_rejects_malformed_adapter_metadata_safely(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _make_document(source)
    block = Block(
        id="block-1",
        seq=0,
        source_text="source",
        source_hash="hash",
        format_metadata={"paragraph_index": True, "style": "Normal"},
    )

    with pytest.raises(DocumentError) as error:
        await docx_adapter.DocxRenderer().render(
            source,
            [block],
            {"block-1": "translation"},
            tmp_path / "out.docx",
        )

    assert error.value.error_code is ErrorCode.RENDER_FAILED
    assert "paragraph index" not in str(error.value)


async def test_renderer_logs_error_class_without_document_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source_text = "PRIVATE SOURCE TEXT"
    translated_text = "PRIVATE TRANSLATED TEXT"

    def fail(*args, **kwargs):
        raise ValueError(f"{source_text}: {translated_text}")

    monkeypatch.setattr(docx_adapter, "_render_sync", fail)
    with capture_logs() as logs, pytest.raises(DocumentError) as error:
        await docx_adapter.DocxRenderer().render(
            tmp_path / "source.docx", [], {}, tmp_path / "out.docx"
        )

    assert error.value.error_code is ErrorCode.RENDER_FAILED
    assert len(logs) == 1
    assert logs[0]["stage"] == "docx_render"
    assert logs[0]["error_type"] == "ValueError"
    assert source_text not in str(logs)
    assert translated_text not in str(logs)
    assert "exc_info" not in logs[0]
