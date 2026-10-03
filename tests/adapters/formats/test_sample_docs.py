"""Generated samples are reproducible and usable through existing format ports."""

import asyncio
from copy import deepcopy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

from app.adapters.formats.docx import DocxExtractor, DocxRenderer
from app.adapters.formats.pdf import PdfExtractor, PdfRenderer
from app.adapters.formats.registry import FormatRegistry
from app.adapters.llm.fake_provider import FakeProvider
from app.core.errors import DocumentError, ErrorCode
from app.core.models import ChunkRequest, TranslationPlan
from scripts.generate_sample_docs import generate_sample_docs


def _docx_translation_details(path: Path) -> tuple[list[str], list[str], list[bytes]]:
    """Include direct cell paragraphs while retaining structure for comparison."""
    document = Document(str(path))
    paragraphs: list[Paragraph] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            paragraphs.append(item)
        else:
            seen: set[Any] = set()
            for row in item.rows:
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        paragraphs.extend(cell.paragraphs)
    table_structure = []
    for table in document.tables:
        element = deepcopy(table._tbl)
        for paragraph in element.xpath("./w:tr/w:tc/w:p"):
            for child in list(paragraph):
                if child.tag != qn("w:pPr"):
                    paragraph.remove(child)
        table_structure.append(etree.tostring(element, method="c14n"))
    return (
        [paragraph.text for paragraph in paragraphs],
        [paragraph.style.name for paragraph in paragraphs],
        table_structure,
    )


async def test_generated_samples_are_reproducible_and_docx_roundtrips(tmp_path: Path) -> None:
    first = await asyncio.to_thread(generate_sample_docs, tmp_path / "first")
    second = await asyncio.to_thread(generate_sample_docs, tmp_path / "second")
    for source, regenerated in zip(first, second, strict=True):
        assert await asyncio.to_thread(source.read_bytes) == await asyncio.to_thread(
            regenerated.read_bytes
        )

    registry = FormatRegistry()
    registry.register("pdf", PdfExtractor(), PdfRenderer())
    registry.register("docx", DocxExtractor(), DocxRenderer())
    source = first[1]
    adapters = await registry.resolve(source)
    assert adapters is not None
    extractor, renderer = adapters
    document = await extractor.extract(source, "sample-docx")
    assert document.blocks[0].source_text == "Document Translation Sample"
    assert document.blocks[-1].source_text == "Project owner: Alex Morgan. Reference: DT-2401."
    result = await FakeProvider(fail_rate=0, latency_ms=0).translate_chunk(
        ChunkRequest(
            chunk_id="sample-docx-chunk",
            blocks=document.blocks,
            target_language="de",
            plan=TranslationPlan(source_language="en", domain="business", register="neutral"),
            model="gpt-4o-mini",
        )
    )
    output = tmp_path / "translated.docx"
    rendered = await renderer.render(source, document.blocks, result.translations, output)
    assert rendered.output_path == output
    assert rendered.degraded_block_ids == []
    original_texts, original_styles, original_tables = await asyncio.to_thread(
        _docx_translation_details, source
    )
    texts, styles, tables = await asyncio.to_thread(_docx_translation_details, output)
    assert styles == original_styles
    assert tables == original_tables
    assert len(tables) == 1
    assert texts == [f"[de] {text.strip()}" if text.strip() else text for text in original_texts]
    assert len(result.translations) == sum(bool(text.strip()) for text in original_texts)


async def test_zip_signature_routes_to_safe_docx_package_validation(tmp_path: Path) -> None:
    malformed = tmp_path / "malformed.docx"
    await asyncio.to_thread(malformed.write_bytes, b"PK\x03\x04not a valid package")
    registry = FormatRegistry()
    registry.register("docx", DocxExtractor(), DocxRenderer())
    adapters = await registry.resolve(malformed)
    assert adapters is not None
    with pytest.raises(DocumentError) as error:
        await adapters[0].extract(malformed, "invalid-document")
    assert error.value.error_code is ErrorCode.CORRUPT_FILE
    assert error.value.retryable is False


@pytest.mark.parametrize("fixture_name", ["platon-gliph", "platon-complex"])
async def test_platon_docx_samples_have_no_glyph_degradation(
    tmp_path: Path, fixture_name: str
) -> None:
    source = Path(__file__).parents[3] / "samples" / f"{fixture_name}.docx"
    document = await DocxExtractor().extract(source, fixture_name)
    translations = {block.id: f"[de] {block.source_text}" for block in document.blocks}

    result = await DocxRenderer().render(
        source, document.blocks, translations, tmp_path / "translated.docx"
    )

    assert document.blocks
    assert document.warnings == []
    assert result.degraded_block_ids == []
    assert result.fallback_blocks == result.fallback_pages == 0
    original_texts, original_styles, original_tables = await asyncio.to_thread(
        _docx_translation_details, source
    )
    texts, styles, tables = await asyncio.to_thread(_docx_translation_details, result.output_path)
    assert texts == [f"[de] {text.strip()}" if text.strip() else text for text in original_texts]
    assert len(translations) == sum(bool(text.strip()) for text in original_texts)
    assert styles == original_styles
    assert tables == original_tables
