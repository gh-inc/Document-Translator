"""Generate deterministic PDF and DOCX fixtures used by format tests."""

from __future__ import annotations

import argparse
import io
import zipfile
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

import pymupdf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

PDF_NAME = "sample_en.pdf"
DOCX_NAME = "sample_en.docx"


def generate_sample_docs(output_dir: Path) -> tuple[Path, Path]:
    """Write the fixed English PDF and DOCX samples to ``output_dir``."""
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = output_dir / PDF_NAME
    docx_path = output_dir / DOCX_NAME
    _generate_pdf(pdf_path)
    _generate_docx(docx_path)
    return pdf_path, docx_path


def _generate_pdf(output_path: Path) -> None:
    document = pymupdf.open()
    first = document.new_page(width=612, height=792)
    first.insert_text((54, 58), "Document Translation Sample", fontsize=20)
    first.insert_text(
        (54, 91),
        "This sample contains clear paragraphs and a compact data table.",
        fontsize=11,
    )
    first.insert_text(
        (54, 111),
        "Keep names, dates, and amounts unchanged when translating this document.",
        fontsize=11,
    )
    first.insert_text((54, 158), "Quarterly Service Summary", fontsize=15)
    _draw_table(
        first,
        (54, 184),
        [
            ("Service", "Requests", "Status"),
            ("Translation", "128", "On track"),
            ("Review", "24", "In progress"),
        ],
    )

    second = document.new_page(width=612, height=792)
    second.insert_text((54, 58), "Delivery Notes", fontsize=18)
    second.insert_text(
        (54, 91),
        "The operations team reviews every completed file before delivery.",
        fontsize=11,
    )
    second.insert_text(
        (54, 111),
        "Please send questions to the project owner by 17 October 2026.",
        fontsize=11,
    )
    second.insert_text((54, 158), "Contact", fontsize=15)
    second.insert_text(
        (54, 183),
        "Project owner: Alex Morgan. Reference: DT-2401.",
        fontsize=11,
    )
    document.set_metadata(
        {
            "title": "Document Translation Sample",
            "author": "Document Translator",
            "subject": "Stable English sample for document adapter tests",
            "keywords": "sample, translation, English",
            "creationDate": "D:20000101000000Z",
            "modDate": "D:20000101000000Z",
        }
    )
    document.save(output_path, garbage=4, deflate=True, no_new_id=True, reproducible=True)
    document.close()


def _draw_table(
    page: pymupdf.Page,
    origin: tuple[float, float],
    rows: list[tuple[str, str, str]],
) -> None:
    x0, y0 = origin
    column_widths = (190.0, 100.0, 160.0)
    row_height = 25.0
    x_positions = [x0]
    for width in column_widths:
        x_positions.append(x_positions[-1] + width)

    for row_index, row in enumerate(rows):
        top = y0 + row_index * row_height
        bottom = top + row_height
        for column_index, value in enumerate(row):
            rectangle = pymupdf.Rect(
                x_positions[column_index],
                top,
                x_positions[column_index + 1],
                bottom,
            )
            page.draw_rect(rectangle, color=(0.2, 0.3, 0.4), width=0.6)
            page.insert_text(
                (rectangle.x0 + 6, rectangle.y0 + 17),
                value,
                fontsize=10,
                fontname="hebo" if row_index == 0 else "helv",
            )


def _generate_docx(output_path: Path) -> None:
    document = Document()
    document.core_properties.title = "Document Translation Sample"
    document.core_properties.subject = "Stable English sample for document adapter tests"
    document.core_properties.author = "Document Translator"
    fixed_time = datetime(2000, 1, 1, 0, 0, 0)
    document.core_properties.created = fixed_time
    document.core_properties.modified = fixed_time

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Document Translation Sample")
    document.add_paragraph(
        "This sample contains styled paragraphs and a compact data table.",
        style="Subtitle",
    )
    document.add_paragraph(
        "Keep names, dates, and amounts unchanged when translating this document.",
        style="Normal",
    )
    document.add_heading("Quarterly Service Summary", level=1)
    _add_docx_table(
        document,
        [
            ("Service", "Requests", "Status"),
            ("Translation", "128", "On track"),
            ("Review", "24", "In progress"),
        ],
    )
    document.add_page_break()
    document.add_heading("Delivery Notes", level=1)
    document.add_paragraph(
        "The operations team reviews every completed file before delivery.",
        style="Normal",
    )
    document.add_paragraph(
        "Please send questions to the project owner by 17 October 2026.",
        style="Normal",
    )
    document.add_heading("Contact", level=2)
    document.add_paragraph(
        "Project owner: Alex Morgan. Reference: DT-2401.",
        style="Normal",
    )
    document.save(str(output_path))
    _normalize_docx_archive(output_path)


def _add_docx_table(document: DocxDocument, rows: list[tuple[str, str, str]]) -> None:
    table = document.add_table(rows=len(rows), cols=3)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value


def _normalize_docx_archive(docx_path: Path) -> None:
    """Fix ZIP member timestamps so repeated generation yields identical bytes."""
    original_bytes = docx_path.read_bytes()
    normalized = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(original_bytes), "r") as source,
        zipfile.ZipFile(normalized, "w", compression=zipfile.ZIP_DEFLATED) as target,
    ):
        for source_info in source.infolist():
            target_info = zipfile.ZipInfo(source_info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            target_info.compress_type = source_info.compress_type
            target_info.external_attr = source_info.external_attr
            target_info.create_system = source_info.create_system
            target_info.flag_bits = source_info.flag_bits
            target.writestr(target_info, source.read(source_info.filename))
    docx_path.write_bytes(normalized.getvalue())


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "samples",
        help="Directory in which to write sample_en.pdf and sample_en.docx.",
    )
    args = parser.parse_args()
    for generated_path in generate_sample_docs(args.output_dir):
        print(generated_path)


if __name__ == "__main__":
    main()
